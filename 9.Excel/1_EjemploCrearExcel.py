import openpyxl

'''
================================================================================
EJERCICIO PRÁCTICO — ACADEMIA DE IDIOMAS "POLYGLOT"
================================================================================

SITUACIÓN PROBLEMA:
  La academia "PolyGlot" necesita un sistema en Python que gestione las
  calificaciones de sus estudiantes. Cada estudiante cursó un idioma (Inglés,
  Francés o Alemán) y tiene tres calificaciones: Parcial1, Parcial2 y Final.

  El sistema debe:
    1. Crear un archivo Excel con los datos de los estudiantes.
    2. Leer ese archivo y calcular el promedio de cada estudiante.
    3. Agregar una columna "Promedio" y otra "Estado" (Aprobado / Reprobado).
    4. Contar cuántos estudiantes aprobaron en cada idioma (estadísticas).
    5. Guardar todo el resultado en un nuevo archivo Excel con dos hojas:
         - "notas"      → tabla completa con promedios y estados
         - "estadisticas" → resumen por idioma

HABILIDADES QUE SE PRACTICAN (basadas en 1_CrearExcel.py):
  ✔ openpyxl.Workbook()         → crear un libro desde cero (nuevo concepto)
  ✔ wb.active / wb.create_sheet → seleccionar / crear hojas
  ✔ ws.append()                 → insertar filas de datos
  ✔ wb.save()                   → guardar en disco
  ✔ openpyxl.load_workbook()    → volver a abrir el archivo guardado
  ✔ ws.iter_rows(values_only)   → recorrer filas con valores directos
  ✔ ws.max_row / ws.max_column  → conocer el tamaño de los datos
  ✔ Diccionario de frecuencias  → contar y acumular valores por categoría
  ✔ Dos hojas en el mismo libro → organizar información relacionada
================================================================================
'''

# ==============================================================================
# RUTAS DE ARCHIVOS
# ==============================================================================

RUTA_BASE   = "/Users/augusto/Documents/FundamentosProgramacion/9.Excel/"
RUTA_DATOS  = RUTA_BASE + "academia_datos.xlsx"    # archivo fuente que crearemos
RUTA_SALIDA = RUTA_BASE + "academia_resultado.xlsx" # archivo final con todo el análisis

HOJA_NOTAS = "notas"
HOJA_EST   = "estadisticas"

# ==============================================================================
# DATOS DE LOS ESTUDIANTES (simulan una base de datos)
# ==============================================================================
# Cada tupla: (Nombre, Idioma, Parcial1, Parcial2, Final)

ESTUDIANTES = [
    ("Ana García",      "Inglés",  85, 90, 88),
    ("Luis Martínez",   "Francés", 60, 55, 70),
    ("Sofía López",     "Alemán",  45, 50, 48),
    ("Pedro Ramírez",   "Inglés",  92, 88, 95),
    ("Valentina Cruz",  "Francés", 78, 82, 80),
    ("Carlos Torres",   "Alemán",  70, 65, 72),
    ("Mariana Díaz",    "Inglés",  55, 60, 58),
    ("Andrés Herrera",  "Francés", 40, 45, 42),
    ("Camila Vargas",   "Alemán",  88, 92, 90),
    ("Jorge Morales",   "Inglés",  73, 77, 75),
]

NOTA_APROBADO = 60   # promedio mínimo para aprobar


# ==============================================================================
# PASO 1 — CREAR EL ARCHIVO EXCEL CON LOS DATOS EN BRUTO
# ==============================================================================
# openpyxl.Workbook() crea un libro de Excel completamente nuevo en memoria
# (equivalente a abrir Excel y hacer "Nuevo archivo").
# El libro viene con una hoja vacía llamada "Sheet" accesible con wb.active.

def crearArchivoFuente(ruta, datos):
    wb = openpyxl.Workbook()   # crea libro vacío en memoria
    ws = wb.active             # obtiene la única hoja que trae por defecto
    ws.title = HOJA_NOTAS      # le cambia el nombre a "notas"

    # Cabeceras de la tabla
    ws.append(["Nombre", "Idioma", "Parcial1", "Parcial2", "Final"])

    # Inserta cada estudiante como una fila
    for estudiante in datos:
        ws.append(list(estudiante))

    wb.save(ruta)
    print(f"[PASO 1] Archivo fuente creado : {ruta}")
    print(f"         Filas de datos        : {ws.max_row - 1}")  # -1 excluye cabecera
    print(f"         Columnas              : {ws.max_column}\n")


# ==============================================================================
# PASO 2 — LEER EL ARCHIVO Y CALCULAR PROMEDIOS
# ==============================================================================
# Carga la hoja "notas" y recorre cada fila de datos (desde fila 2 para saltar
# la cabecera) usando iter_rows con values_only=True.
# Calcula el promedio de las tres notas y asigna estado "Aprobado" o "Reprobado".

def calcularPromedios(ruta_in, hoja):
    wb = openpyxl.load_workbook(ruta_in)
    ws = wb[hoja]

    print(f"[PASO 2] Leyendo '{hoja}' desde : {ruta_in}")
    print(f"         Dimensiones de la hoja : {ws.dimensions}\n")

    resultados = []   # acumula filas procesadas: (nombre, idioma, p1, p2, final, promedio, estado)

    # iter_rows(min_row=2) salta la cabecera; values_only devuelve valores directos (no objetos Cell)
    for fila in ws.iter_rows(min_row=2, values_only=True):
        nombre, idioma, p1, p2, final = fila          # desempaqueta los 5 valores de la fila
        promedio = round((p1 + p2 + final) / 3, 1)   # promedio con 1 decimal
        estado   = "Aprobado" if promedio >= NOTA_APROBADO else "Reprobado"
        resultados.append((nombre, idioma, p1, p2, final, promedio, estado))

    return resultados


# ==============================================================================
# PASO 3 — CONTAR APROBADOS POR IDIOMA (DICCIONARIO DE FRECUENCIAS)
# ==============================================================================
# Recorre los resultados ya calculados y construye un diccionario con la
# siguiente estructura:
#   { "Inglés": {"total": 4, "aprobados": 3}, "Francés": {...}, ... }

def contarAprobadosPorIdioma(resultados):
    estadisticas = {}

    for (_, idioma, _, _, _, _, estado) in resultados:   # _ descarta valores que no usamos
        if idioma not in estadisticas:
            estadisticas[idioma] = {"total": 0, "aprobados": 0}

        estadisticas[idioma]["total"] += 1
        if estado == "Aprobado":
            estadisticas[idioma]["aprobados"] += 1

    print("[PASO 3] Estadísticas por idioma:")
    for idioma, datos in estadisticas.items():
        print(f"         {idioma:<10} → total: {datos['total']}  |  aprobados: {datos['aprobados']}")
    print()

    return estadisticas


# ==============================================================================
# PASO 4 — GUARDAR EL RESULTADO EN UN ARCHIVO NUEVO CON DOS HOJAS
# ==============================================================================
# Crea un libro nuevo con dos hojas:
#   "notas"        → tabla completa con promedios y estados
#   "estadisticas" → resumen de aprobados por idioma

def guardarResultado(ruta_out, resultados, estadisticas):
    wb = openpyxl.Workbook()

    # ---------- HOJA 1: notas completas ----------
    ws_notas = wb.active
    ws_notas.title = HOJA_NOTAS
    ws_notas.append(["Nombre", "Idioma", "Parcial1", "Parcial2", "Final", "Promedio", "Estado"])

    for fila in resultados:
        ws_notas.append(list(fila))

    # ---------- HOJA 2: estadísticas por idioma ----------
    # wb.create_sheet() agrega una nueva hoja al libro
    ws_est = wb.create_sheet(HOJA_EST)
    ws_est.append(["Idioma", "Total Estudiantes", "Aprobados", "Reprobados", "% Aprobación"])

    for idioma, datos in estadisticas.items():
        total      = datos["total"]
        aprobados  = datos["aprobados"]
        reprobados = total - aprobados
        porcentaje = round((aprobados / total) * 100, 1)
        ws_est.append([idioma, total, aprobados, reprobados, porcentaje])

    wb.save(ruta_out)
    print(f"[PASO 4] Resultado guardado en : {ruta_out}")
    print(f"         Hojas creadas         : {wb.sheetnames}\n")


# ==============================================================================
# PASO 5 — VERIFICAR EL RESULTADO LEYENDO EL ARCHIVO GUARDADO
# ==============================================================================
# Abre el archivo de salida y muestra ambas hojas como matrices para verificar
# que todo quedó guardado correctamente.

def verificarResultado(ruta_out):
    print("[PASO 5] Verificando contenido del archivo guardado...\n")

    wb = openpyxl.load_workbook(ruta_out)

    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        print(f"  --- Hoja: '{nombre_hoja}' ({ws.max_row} filas × {ws.max_column} columnas) ---")

        for fila in ws.iter_rows(values_only=True):
            print(f"    {list(fila)}")
        print()


# ==============================================================================
# PROGRAMA PRINCIPAL
# ==============================================================================

if __name__ == "__main__":
    print("=" * 60)
    print("  SISTEMA DE GESTIÓN — ACADEMIA POLYGLOT")
    print("=" * 60 + "\n")

    crearArchivoFuente(RUTA_DATOS, ESTUDIANTES)        # Paso 1
    resultados    = calcularPromedios(RUTA_DATOS, HOJA_NOTAS)   # Paso 2
    estadisticas  = contarAprobadosPorIdioma(resultados)         # Paso 3
    guardarResultado(RUTA_SALIDA, resultados, estadisticas)      # Paso 4
    verificarResultado(RUTA_SALIDA)                              # Paso 5

    print("=" * 60)
    print("  PROCESO COMPLETADO EXITOSAMENTE")
    print("=" * 60)
